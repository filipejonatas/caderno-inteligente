from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
from openpyxl import load_workbook as open_xlsx

from caderno_inteligente.ingestion import load_workbook
from caderno_inteligente.schemas import SCHEMAS
from caderno_inteligente.workbook_writer import EXAMPLE_MARKER, write_workbook

SOURCE = Path("data/source/Base de Dados - Caderno Inteligente.xlsm")
TITLE = "Teste do gravador"


def _read(content, sheet):
    return pd.read_excel(BytesIO(content), sheet_name=sheet, skiprows=2)


def test_roundtrip_keeps_values_in_the_source_layout():
    data = load_workbook(SOURCE)
    content = write_workbook(data, title=TITLE)
    for sheet in ("Vendas_24m", "Carteira_Pedidos", "Capacidade_Semanal"):
        pd.testing.assert_frame_equal(_read(content, sheet), data[sheet], check_dtype=False)
    sheet = open_xlsx(BytesIO(content))["Vendas_24m"]
    assert TITLE in sheet["A1"].value and sheet["A2"].value is None
    assert sheet["A3"].value == "Mês" and sheet["A4"].value == "2024-09"


def test_sheet_order_notes_and_selection():
    data = load_workbook(SOURCE)
    content = write_workbook(data, title=TITLE, notes=[f"{EXAMPLE_MARKER} - nota"])
    assert open_xlsx(BytesIO(content)).sheetnames == ["LEIA_ME", *SCHEMAS, "Precos_Produtos"]
    only = write_workbook(data, title=TITLE, sheets=["Estoque_Atual", "Produtos"])
    assert open_xlsx(BytesIO(only)).sheetnames == ["Produtos", "Estoque_Atual"]


def test_dates_are_formatted_and_no_macro_or_formula():
    data = load_workbook(SOURCE)
    data["Produtos"].loc[0, "Produto"] = "=1+1"
    content = write_workbook(data, title=TITLE)
    book = open_xlsx(BytesIO(content))
    assert book["Capacidade_Semanal"]["A4"].number_format == "dd/mm/yyyy"
    assert book["Produtos"]["B4"].data_type == "s" and book["Produtos"]["B4"].value == "=1+1"
    assert not any(name.endswith("vbaProject.bin") for name in ZipFile(BytesIO(content)).namelist())
    assert book.properties.creator == "Caderno Inteligente"
    assert book.properties.created.isoformat().startswith("2026-10-12T00:00:00")


def test_missing_values_stay_empty_not_zero():
    frame = pd.DataFrame({"SKU": ["A", "B"], "Estoque atual": [1.0, float("nan")]})
    book = open_xlsx(BytesIO(write_workbook({"Produtos": frame}, title=TITLE)))
    assert book["Produtos"]["B5"].value is None
