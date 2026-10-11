"""Rotas /api/dataset (U1A) num app FastAPI próprio: SQLite em tmp_path, motor e informações da base falsos."""
import hashlib
import io
import json
import logging
import time
import zipfile
from pathlib import Path

import pytest
from fastapi import Depends, FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook as open_xlsx

from backend.dataset import create_dataset_router
from caderno_inteligente.auth import UserStore, issue_token
from caderno_inteligente.dataset_store import Database, DatasetStore, DatasetStoreError, content_sha256
from caderno_inteligente.dataset_upload import reference_of
from caderno_inteligente.ingestion import load_workbook
from caderno_inteligente.transformations import safe_normalise
from caderno_inteligente.upload_safety import DEFAULT_PATH
from caderno_inteligente.workbook_writer import write_workbook

SOURCE = Path("data/source/Base de Dados - Caderno Inteligente.xlsm")
EXAMPLE = Path("docs/exemplos/atualizacao-2026-10.xlsx")
SECRET = "segredo-de-teste-" + "x" * 30
EMAIL = "pcp@exemplo.com"
TIMINGS: dict[str, float] = {}


@pytest.fixture(scope="module")
def original():
    return load_workbook(SOURCE)


class Env:
    def __init__(self, tmp_path, original, monkeypatch, *, settings_file=DEFAULT_PATH, supply=None):
        monkeypatch.setenv("AUTH_SECRET", SECRET)
        monkeypatch.delenv("AUTH_REQUIRED", raising=False)
        self.source = "banco"
        self.writable = True
        self.changes = 0
        self.info_override: dict = {}
        self.database = Database(sqlite_path=tmp_path / "dataset.db")
        self.store = DatasetStore(self.database)
        self.store.import_dataset(original)
        UserStore(self.database).create_user(EMAIL, "Ana PCP", "senha-gerada-no-teste")

        def write_guard():
            if not self.writable:
                raise HTTPException(403, "Somente leitura.")

        def info():
            dataset = self.store.load_dataset() if self.source == "banco" else original
            reference, sales = reference_of(safe_normalise(dataset)[0])
            return {"reference_date": reference, "reference_source": "auto", "sales_last_month": sales,
                    "content_sha256": content_sha256(dataset), "source_sha256": "f" * 64,
                    "matches_original": content_sha256(dataset) == content_sha256(original), **self.info_override}

        def engine(dataset):
            return {"queue_size": 0 if dataset is None else len(dataset["Produtos"]), "top5": []}

        def changed():
            self.changes += 1

        app = FastAPI()
        app.include_router(create_dataset_router(
            production=lambda: False, data_source=lambda: self.source, store=lambda: self.store,
            users=lambda: UserStore(self.database), write_access=Depends(write_guard), settings_file=settings_file,
            current_info=info, engine_summary=engine, original_loader=lambda: original, on_change=changed,
            describe_error=lambda message, error: message,
            supply_settings=None if supply is None else (lambda: dict(supply)),
        ))
        self.client = TestClient(app)
        self.auth = {"Authorization": "Bearer " + issue_token(SECRET, EMAIL, "Ana PCP", 3600)[0]}

    def upload(self, path: str, content: bytes, *, filename="envio.xlsx", headers=None, auth=True):
        sent = {"Content-Type": "application/octet-stream", "X-Filename": filename, **(self.auth if auth else {}), **(headers or {})}
        return self.client.post(f"/api/dataset/{path}", content=content, headers=sent)

    def apply(self, content: bytes, *, version=None, confirm=None, **kwargs):
        headers = {"X-Confirm-Sha256": confirm or hashlib.sha256(content).hexdigest(),
                   "X-Expected-Version": str(self.store.version() if version is None else version)}
        return self.upload("apply", content, headers=headers, **kwargs)

    def post(self, path: str, *, version=None, **kwargs):
        """Voltar para versão e restaurar a original exigem X-Expected-Version (padrão: a versão atual)."""
        headers = {**self.auth, "X-Expected-Version": str(self.store.version() if version is None else version)}
        return self.client.post(f"/api/dataset/{path}", headers=headers, **kwargs)


@pytest.fixture
def env(tmp_path, original, monkeypatch):
    return Env(tmp_path, original, monkeypatch)


def test_preview_has_no_effect_and_apply_writes_a_new_version(env, caplog):
    """Critério 6 e o pacote de exemplo pela API: a prévia não muda a versão; aplicar grava, avisa o cache e loga sem e-mail."""
    content = EXAMPLE.read_bytes()
    before = env.store.version()
    caplog.set_level(logging.INFO, logger="caderno_inteligente.api")
    started = time.perf_counter()
    preview = env.upload("preview", content, filename="atualizacao-2026-10.xlsx")
    TIMINGS["preview"] = time.perf_counter() - started
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["can_apply"] is True and body["expected_version"] == before
    assert body["reference"]["after"] == "2026-10-12" and body["file"]["example"] is True
    assert env.store.version() == before and env.changes == 0

    started = time.perf_counter()
    applied = env.apply(content, filename="atualizacao-2026-10.xlsx")
    TIMINGS["apply"] = time.perf_counter() - started
    assert applied.status_code == 200, applied.text
    result = applied.json()
    assert result["version"] == before + 1 and result["reference_date"] == "2026-10-12" and result["sales_last_month"] == "2026-09"
    assert result["content_sha256"] == content_sha256(env.store.load_dataset()) and env.changes == 1
    [row] = env.store.list_versions()
    assert (row["action"], row["filename"], row["user_email"]) == ("upload", "atualizacao-2026-10.xlsx", EMAIL)
    assert row["summary"]["engine"]["before"] is None and row["summary"]["can_apply"] is True
    assert "dataset_upload action=preview" in caplog.text and "dataset_upload action=apply" in caplog.text
    assert EMAIL not in caplog.text
    print(f"\ntempo SQLite: prévia {TIMINGS['preview']:.2f} s, aplicação {TIMINGS['apply']:.2f} s")


def test_versions_round_trip_keeps_the_content_hash(env, original):
    """Critério 9: aplicar, voltar para a versão anterior e restaurar a original devolvem os mesmos content_sha256."""
    first = env.post("restore-original")
    assert first.status_code == 200, first.text
    original_sha = first.json()["content_sha256"]
    assert original_sha == content_sha256(original)
    applied = env.apply(EXAMPLE.read_bytes())
    assert applied.status_code == 200, applied.text
    uploaded_sha = applied.json()["content_sha256"]
    assert uploaded_sha != original_sha

    back = env.post(f"versions/{first.json()['version_id']}/restore")
    assert back.status_code == 200, back.text
    assert back.json()["content_sha256"] == original_sha and back.json()["reference_date"] == "2026-09-14"
    forward = env.post(f"versions/{applied.json()['version_id']}/restore")
    assert forward.json()["content_sha256"] == uploaded_sha
    again = env.post("restore-original")
    assert again.json()["content_sha256"] == original_sha

    listing = env.client.get("/api/dataset/versions").json()
    assert [item["action"] for item in listing["items"]] == ["original", "restore", "restore", "upload", "original"]
    assert [item["active"] for item in listing["items"]] == [True, False, False, False, False]
    assert "summary" not in listing["items"][0] and listing["active_version_id"] == again.json()["version_id"]
    assert env.post("versions/99999/restore").status_code == 404

    status = env.client.get("/api/dataset", headers=env.auth).json()
    assert (status["origin"], status["matches_original"], status["upload_enabled"]) == ("original", True, True)
    assert status["version"] == again.json()["version"] and status["uploaded_by"] == EMAIL


def test_protections_and_file_rejections(env):
    """Critério 8: 413 (com e sem Content-Length), ZIP grande e CSV renomeado 422, 403, 409 e 401, nunca 500."""
    big = b"PK\x03\x04" + b"0" * (4 * 1024 * 1024 + 10)
    assert env.upload("preview", big).status_code == 413
    chunked = env.client.post("/api/dataset/preview", content=(big[i:i + 65536] for i in range(0, len(big), 65536)),
                              headers={**env.auth, "X-Filename": "envio.xlsx"})
    assert chunked.status_code == 413 and chunked.json()["code"] == "FILE_TOO_LARGE"

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("xl/worksheets/sheet1.xml", b"\0" * (60 * 1024 * 1024))
    bomb = env.upload("preview", buffer.getvalue())
    assert bomb.status_code == 422 and bomb.json()["code"].startswith("ZIP_")

    renamed = env.upload("preview", b"SKU;Estoque atual\nCI-0001;10\n", filename="dados.xlsx")
    assert renamed.status_code == 422 and renamed.json()["code"] == "NOT_XLSX"
    assert env.upload("preview", EXAMPLE.read_bytes(), auth=False).status_code == 401
    assert env.upload("preview", EXAMPLE.read_bytes(), headers={"Authorization": env.auth["Authorization"][:-3] + "abc"}).status_code == 401
    env.writable = False
    assert env.upload("preview", EXAMPLE.read_bytes()).status_code == 403
    assert env.post("restore-original").status_code == 403
    env.writable, env.source = True, "planilha"
    assert env.upload("preview", EXAMPLE.read_bytes()).status_code == 409
    assert env.post("restore-original").status_code == 409


def test_apply_checks_the_confirmation_and_the_expected_version(env):
    content = EXAMPLE.read_bytes()
    version = env.store.version()
    mismatch = env.apply(content, confirm="0" * 64)
    assert mismatch.status_code == 422 and mismatch.json()["code"] == "CONFIRM_MISMATCH"
    stale = env.apply(content, version=version - 1)
    assert stale.status_code == 409 and stale.json()["detail"] == "Outra pessoa atualizou a base; refaça a prévia."
    missing = env.upload("apply", content, headers={"X-Confirm-Sha256": hashlib.sha256(content).hexdigest()})
    assert missing.status_code == 422
    assert env.store.version() == version and env.store.list_versions() == []


def test_data_error_returns_422_with_the_address_and_the_preview(env, original):
    """Critério 7 pela API: erro de dado → 422 com o endereço, e a aplicação também recusa."""
    stock = original["Estoque_Atual"].copy()
    stock["Estoque atual"] = stock["Estoque atual"].astype(object)
    stock.loc[0, "Estoque atual"] = "muito"
    content = write_workbook({"Estoque_Atual": stock}, title="teste")
    response = env.upload("preview", content)
    assert response.status_code == 422
    body = response.json()
    assert body["detail"].startswith("A planilha tem ") and body["preview"]["can_apply"] is False
    error = next(item for item in body["preview"]["errors"] if item["code"] == "INVALID_NUMBER")
    assert (error["sheet"], error["row"], error["value"]) == ("Estoque_Atual", 4, "muito") and error["cell"].endswith("4")
    assert error["message"] == f"Estoque_Atual!{error['cell']}: 'muito' não é número"
    assert env.apply(content).status_code == 422 and env.store.list_versions() == []


def test_status_reports_why_upload_is_disabled(env):
    env.source = "planilha"
    body = env.client.get("/api/dataset").json()
    assert (body["upload_enabled"], body["upload_disabled_reason"], body["origin"], body["version"]) == (False, "planilha", "original", None)
    assert body["max_upload_bytes"] == 4 * 1024 * 1024 and body["accepted_extensions"] == [".xlsx", ".xlsm"]
    assert body["reference_date"] == "2026-09-14" and body["upload_disabled_message"]
    assert env.client.get("/api/dataset/versions").json()["items"] == []
    env.source = "banco"
    body = env.client.get("/api/dataset").json()
    assert (body["upload_disabled_reason"], body["origin"], body["active_version_id"]) == ("login", "original", None)
    env.writable = False
    assert env.client.get("/api/dataset", headers=env.auth).json()["upload_disabled_reason"] == "somente_leitura"
    env.writable = True
    assert env.client.get("/api/dataset", headers=env.auth).json()["upload_enabled"] is True


def test_template_has_the_thirteen_sheets_with_header_on_row_three(env, original):
    response = env.client.get("/api/dataset/template")
    assert response.status_code == 200
    assert response.headers["content-disposition"] == 'attachment; filename="modelo-caderno-inteligente-2026-09-14.xlsx"'
    workbook = open_xlsx(io.BytesIO(response.content), read_only=True)
    assert len(workbook.sheetnames) == 13 and set(workbook.sheetnames) == set(original)
    sheet = workbook["Produtos"]
    assert [cell.value for cell in next(sheet.iter_rows(min_row=3, max_row=3))] == list(original["Produtos"].columns)
    assert "14/09/2026" in sheet["A1"].value


def test_example_download_is_the_fictional_package(env):
    response = env.client.get("/api/dataset/example")
    assert response.status_code == 200
    assert response.headers["content-disposition"] == 'attachment; filename="atualizacao-2026-10.xlsx"'
    assert "LEIA_ME" in open_xlsx(io.BytesIO(response.content), read_only=True).sheetnames


def test_restore_routes_require_the_expected_version(env):
    """Voltar para versão e restaurar a original: 422 sem X-Expected-Version, 409 com versão desatualizada; nada é gravado."""
    first = env.post("restore-original")
    assert first.status_code == 200, first.text
    version_id, version = first.json()["version_id"], env.store.version()
    for path in ("restore-original", f"versions/{version_id}/restore"):
        missing = env.client.post(f"/api/dataset/{path}", headers=env.auth)
        assert missing.status_code == 422 and "X-Expected-Version" in missing.json()["detail"]
        stale = env.post(path, version=version - 1)
        assert stale.status_code == 409 and stale.json()["detail"] == "Outra pessoa atualizou a base; refaça a prévia."
    assert env.store.version() == version and len(env.store.list_versions()) == 1


def test_uploader_email_is_masked_for_anonymous_viewers(env):
    assert env.post("restore-original").status_code == 200
    anonymous = env.client.get("/api/dataset").json()
    assert anonymous["uploaded_by"] == "p***@exemplo.com"
    assert [item["user_email"] for item in env.client.get("/api/dataset/versions").json()["items"]] == ["p***@exemplo.com"]
    assert env.client.get("/api/dataset", headers=env.auth).json()["uploaded_by"] == EMAIL
    assert env.client.get("/api/dataset/versions", headers=env.auth).json()["items"][0]["user_email"] == EMAIL


def test_reference_fixed_in_config_is_used_by_the_preview_and_the_saved_version(tmp_path, original, monkeypatch):
    """Data fixa em config/supply_plan.json: a prévia bloqueia vendas além dela e a versão grava a mesma data do motor."""
    env = Env(tmp_path, original, monkeypatch, supply={"reference_date": "2026-09-14"})
    preview = env.upload("preview", EXAMPLE.read_bytes())
    assert preview.status_code == 422
    body = preview.json()["preview"]
    assert body["reference"]["after"] == "2026-09-14"
    assert "REFERENCE_FIXED_IN_CONFIG" in {item["code"] for item in body["errors"]}

    later = Env(tmp_path / "later", original, monkeypatch, supply={"reference_date": "2026-09-21"})
    saved = later.post("restore-original")
    assert saved.status_code == 200, saved.text
    assert saved.json()["reference_date"] == "2026-09-21" and later.store.list_versions()[0]["reference_date"] == "2026-09-21"


def test_template_keeps_inactive_skus_and_the_preview_reference_comes_from_the_fresh_base(env, original):
    env.store.set_active("CI-0050", False, EMAIL)
    workbook = open_xlsx(io.BytesIO(env.client.get("/api/dataset/template").content), read_only=True)
    skus = [row[0] for row in workbook["Produtos"].iter_rows(min_row=4, values_only=True)]
    assert "CI-0050" in skus and len(skus) == len(original["Produtos"])
    # A referência "antes" vem do conjunto carregado agora, não do current_info (que poderia estar em cache).
    env.info_override = {"reference_date": "2000-01-03", "sales_last_month": "1999-12"}
    body = env.upload("preview", EXAMPLE.read_bytes()).json()
    assert (body["reference"]["before"], body["reference"]["sales_last_month_before"]) == ("2026-09-14", "2026-08")


def test_saved_versions_follow_max_versions_from_the_settings(tmp_path, original, monkeypatch):
    settings = json.loads(DEFAULT_PATH.read_text(encoding="utf-8"))
    path = tmp_path / "dataset_upload.json"
    path.write_text(json.dumps({**settings, "max_versions": 2}), encoding="utf-8")
    env = Env(tmp_path, original, monkeypatch, settings_file=path)
    for _ in range(3):
        assert env.post("restore-original").status_code == 200
    assert len(env.store.list_versions()) == 2


def test_missing_migration_returns_503_with_the_store_message(env, monkeypatch):
    message = "O banco não tem a tabela de versões; rode a migração 005."

    def broken():
        raise DatasetStoreError(message)

    monkeypatch.setattr(env.store, "state", broken)
    for response in (env.client.get("/api/dataset"), env.client.get("/api/dataset/versions"),
                     env.upload("preview", EXAMPLE.read_bytes()), env.post("restore-original"), env.apply(EXAMPLE.read_bytes())):
        assert response.status_code == 503 and response.json()["detail"] == message
