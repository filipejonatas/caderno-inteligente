from pathlib import Path
import pytest
from caderno_inteligente.ingestion import IngestionError, load_workbook

SOURCE = Path("data/source/Base de Dados - Caderno Inteligente.xlsm")

def test_loads_required_sheets_without_mutating_source():
    before = SOURCE.stat().st_mtime_ns
    data = load_workbook(SOURCE)
    assert {"Produtos", "Sell_Out", "Capacidade_Semanal"}.issubset(data)
    assert len(data["Produtos"]) == 50
    assert SOURCE.stat().st_mtime_ns == before

def test_rejects_missing_file():
    with pytest.raises(IngestionError, match="não encontrado"):
        load_workbook("inexistente.xlsm")

def _workbook_bytes(sheets: dict[str, list[list]]) -> bytes:
    from io import BytesIO
    import openpyxl
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()

def test_read_workbook_bytes_matches_load_workbook_on_packaged_source():
    import pandas as pd
    from caderno_inteligente.ingestion import read_workbook_bytes
    expected = load_workbook(SOURCE)
    result = read_workbook_bytes(SOURCE.read_bytes())
    assert set(result["sheets"]) == set(expected) and len(expected) == 13
    for sheet, frame in expected.items():
        pd.testing.assert_frame_equal(result["sheets"][sheet], frame)
    assert result["header_rows"] == {sheet: 3 for sheet in expected}
    assert result["missing_header"] == {}
    assert "LEIA_ME" in result["ignored_sheets"] and "Dicionario_Dados" in result["ignored_sheets"]
    assert result["example_marker"] is False

def test_read_workbook_bytes_detects_clean_header_missing_header_and_example_marker():
    from caderno_inteligente.ingestion import read_workbook_bytes
    content = _workbook_bytes({
        "Lead_Times": [["SKU", "Família", "Lead time dias", "Lote mínimo"], ["CI-0001", "Escolar", 10, 100]],
        "Estoque_Atual": [["SKU", "Estoque atual"], ["CI-0001", 5]],
        "LEIA_ME": [["Natureza", "Arquivo de exemplo fictício para testar o botão"]],
    })
    result = read_workbook_bytes(content)
    assert result["header_rows"] == {"Lead_Times": 1}
    assert result["sheets"]["Lead_Times"].loc[0, "Lote mínimo"] == 100
    assert result["missing_header"] == {"Estoque_Atual": ["Cobertura dias", "Estoque segurança dias"]}
    assert result["ignored_sheets"] == ["LEIA_ME"]
    assert result["example_marker"] is True

def test_read_workbook_bytes_rejects_unreadable_content():
    from caderno_inteligente.ingestion import read_workbook_bytes
    from caderno_inteligente.upload_safety import UploadRejected
    with pytest.raises(UploadRejected) as error:
        read_workbook_bytes(b"isto nao e uma planilha")
    assert error.value.code == "NOT_XLSX" and error.value.status == 422

def test_read_workbook_bytes_corrupted_inside_is_not_xlsx_never_500():
    """sheetId não numérico (TypeError) e atributo duplicado (ParseError) viram 422 NOT_XLSX."""
    from io import BytesIO
    from zipfile import ZipFile
    from caderno_inteligente.ingestion import read_workbook_bytes
    from caderno_inteligente.upload_safety import UploadRejected
    base = _workbook_bytes({"Lead_Times": [["SKU", "Família", "Lead time dias", "Lote mínimo"], ["CI-0001", "Escolar", 10, 100]]})
    workbook_xml = ZipFile(BytesIO(base)).read("xl/workbook.xml")
    for broken in (workbook_xml.replace(b'sheetId="1"', b'sheetId="abc"'), workbook_xml.replace(b'sheetId="1"', b'sheetId="1" sheetId="2"')):
        buffer = BytesIO()
        with ZipFile(BytesIO(base)) as source, ZipFile(buffer, "w") as target:
            for name in source.namelist():
                target.writestr(name, broken if name == "xl/workbook.xml" else source.read(name))
        with pytest.raises(UploadRejected) as error:
            read_workbook_bytes(buffer.getvalue())
        assert (error.value.status, error.value.code) == (422, "NOT_XLSX")

def test_read_workbook_bytes_reads_at_most_max_rows_plus_one():
    from caderno_inteligente.ingestion import read_workbook_bytes
    rows = [["SKU", "Família", "Lead time dias", "Lote mínimo"]] + [[f"CI-{index:04d}", "Escolar", 10, 100] for index in range(30)]
    result = read_workbook_bytes(_workbook_bytes({"Lead_Times": rows}), max_rows=10)
    assert len(result["sheets"]["Lead_Times"]) == 11  # Uma além do limite: a prévia ainda aponta TOO_MANY_ROWS.
