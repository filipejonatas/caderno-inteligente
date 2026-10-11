from __future__ import annotations
from datetime import datetime
from io import BytesIO
from typing import Iterable
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from .schemas import SCHEMAS

EXAMPLE_MARKER = "EXEMPLO FICTÍCIO"
README_SHEET = "LEIA_ME"
SHEET_ORDER = (*SCHEMAS, "Precos_Produtos")
HEADER_ROW = 3  # A1 = título, linha 2 vazia, cabeçalho na 3 e dados a partir da 4 (mesmo layout da planilha de origem).
DATE_FORMAT = "dd/mm/yyyy"
FIXED_TIMESTAMP = datetime(2026, 10, 12)
CREATOR = "Caderno Inteligente"
MAX_COLUMN_WIDTH = 60


def _cell_value(value):
    """Converte para tipos nativos; dado ausente vira célula vazia, nunca zero."""
    if value is None or value is pd.NaT:
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item") and not isinstance(value, (str, bytes)):  # escalares do numpy
        value = value.item()
    if isinstance(value, float) and value != value:
        return None
    return value


def _put(sheet, row: int, column: int, value):
    cell = sheet.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # texto que começa com "=" não pode virar fórmula
    elif isinstance(value, datetime):
        cell.number_format = DATE_FORMAT
    return cell


def _write_sheet(workbook: Workbook, name: str, frame: pd.DataFrame, title: str) -> None:
    sheet = workbook.create_sheet(name)
    sheet.cell(row=1, column=1, value=f"{title} · {name}").font = Font(bold=True, size=13)
    widths: list[int] = []
    for position, column in enumerate(frame.columns, start=1):
        _put(sheet, HEADER_ROW, position, str(column)).font = Font(bold=True)
        widths.append(len(str(column)))
    for offset, record in enumerate(frame.itertuples(index=False, name=None)):
        for position, raw in enumerate(record, start=1):
            value = _cell_value(raw)
            _put(sheet, HEADER_ROW + 1 + offset, position, value)
            shown = value.strftime("%d/%m/%Y") if isinstance(value, datetime) else ("" if value is None else str(value))
            widths[position - 1] = max(widths[position - 1], len(shown))
    for position, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = min(width + 2, MAX_COLUMN_WIDTH)
    sheet.freeze_panes = sheet.cell(row=HEADER_ROW + 1, column=1)


def write_workbook(dataset: dict[str, pd.DataFrame], *, title: str, notes: list[str] | None = None,
                   sheets: Iterable[str] | None = None) -> bytes:
    """Grava um .xlsx sem macro nem fórmula, no layout lido por read_workbook_bytes (cabeçalho na linha 3)."""
    wanted = None if sheets is None else set(sheets)
    workbook = Workbook()
    workbook.remove(workbook.active)
    if notes:
        readme = workbook.create_sheet(README_SHEET)
        readme.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
        for offset, note in enumerate(notes):
            _put(readme, HEADER_ROW + offset, 1, str(note))
        readme.column_dimensions["A"].width = MAX_COLUMN_WIDTH + 40
    for name in SHEET_ORDER:
        if name in dataset and (wanted is None or name in wanted):
            _write_sheet(workbook, name, dataset[name], title)
    if not workbook.sheetnames:
        workbook.create_sheet(README_SHEET)
    properties = workbook.properties
    properties.creator = CREATOR
    properties.lastModifiedBy = CREATOR
    properties.created = FIXED_TIMESTAMP
    properties.modified = FIXED_TIMESTAMP
    properties.title = title
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
